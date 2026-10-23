"""Arma el libro Excel (lenguaje de formatos BAC-14) a partir de una carpeta de CSV, uno por hoja.

Uso: python3 scripts/csv_to_xlsx.py <carpeta_csv> <salida.xlsx>
"""
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ORDEN = [
    "00_GUIA", "Sistema", "ParticipanteExterno", "Area", "Intercambio", "GrupoInformacion",
    "EstadoPedido", "ModalidadEntrega", "Problema", "Fuente",
    "RutaExpress_marketplaces", "RutaExpress_aliados", "Intercambio_transporta",
    "Intercambio_estadosReportados", "Intercambio_modalidades", "Intercambio_problemas",
    "Intercambio_fuentes", "Problema_reportadoPor", "Problema_afectados", "Problema_fuentes",
    "EstadoPedido_siguientes", "Catalogos", "POR_LEVANTAR",
]
# Columnas que toman su lista de valores de la hoja Catalogos.
ENUMS = {"mecanismo": "Mecanismo", "oportunidad": "Oportunidad", "tipo": "TipoFuente", "Nivel": "Nivel"}

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def leer(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        return [row for row in csv.reader(f)]


def main(carpeta, salida):
    carpeta = Path(carpeta)
    hojas = {p.stem: leer(p) for p in carpeta.glob("*.csv")}
    nombres = [n for n in ORDEN if n in hojas] + sorted(n for n in hojas if n not in ORDEN)

    wb = Workbook()
    wb.remove(wb.active)
    rangos_enum = {}
    if "Catalogos" in hojas:
        filas = hojas["Catalogos"][1:]
        for enum in {r[0] for r in filas if r}:
            idx = [i + 2 for i, r in enumerate(filas) if r and r[0] == enum]
            rangos_enum[enum] = f"Catalogos!$B${min(idx)}:$B${max(idx)}"

    for nombre in nombres:
        filas = hojas[nombre]
        ws = wb.create_sheet(nombre[:31])
        for fila in filas:
            ws.append(fila)
        for celda in ws[1]:
            celda.fill, celda.font = HEADER_FILL, HEADER_FONT
            celda.alignment = Alignment(wrap_text=True, vertical="center")
        for col in ws.columns:
            largo = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(12, largo + 2), 60)
            for c in col[1:]:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        for celda in ws[1]:
            enum = ENUMS.get(celda.value)
            if enum and enum in rangos_enum and nombre != "Catalogos":
                dv = DataValidation(type="list", formula1=rangos_enum[enum], allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f"{celda.column_letter}2:{celda.column_letter}500")
    wb.save(salida)
    print(f"{salida}: {len(nombres)} hojas")


if __name__ == "__main__":
    main(*sys.argv[1:3])
